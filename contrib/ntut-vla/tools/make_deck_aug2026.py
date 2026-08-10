"""
Build the August 2026 progress deck, in the same house style as the July one.

Every number here comes from a logged flight or a logged measurement, and the
sources are named on the slides so any figure can be traced back to a file.

Run:  python tools/make_deck_aug2026.py
Out:  docs/VLA-Guardrail-Follow-Aug2026.pptx
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "VLA-Guardrail-Follow-Aug2026.pptx"

FONT = "Poppins"
TEAL = RGBColor(0x24, 0x9D, 0xB2)
TEAL_DK = RGBColor(0x1A, 0x74, 0x84)
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x2B, 0x2B, 0x2B)
GREY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xEA, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF4, 0xF7, 0xF9)
GOOD = RGBColor(0x1B, 0x8A, 0x5A)
BAD = RGBColor(0xC0, 0x43, 0x3F)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
BLANK = prs.slide_layouts[6]
_page = [0]


def _txt(slide, x, y, w, h, text, size=11, bold=False, colour=BODY,
         align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
    return box


def _rect(slide, x, y, w, h, fill=CARD, line=None, radius=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def slide(eyebrow, title, sub=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    _txt(s, 0.55, 0.42, 8.5, 0.28, eyebrow.upper(), 10, True, TEAL)
    _txt(s, 0.55, 0.72, 8.9, 0.62, title, 25, True, INK, spacing=0.95)
    y = 1.44
    if sub:
        _txt(s, 0.55, 1.40, 8.9, 0.34, sub, 11.5, False, GREY)
        y = 1.85
    _page[0] += 1
    _rect(s, 9.30, 5.18, 0.50, 0.32, TEAL)
    _txt(s, 9.30, 5.21, 0.50, 0.28, f"{_page[0]:02d}", 9, True, WHITE,
         align=PP_ALIGN.CENTER)
    return s, y


def kpi(s, x, y, w, value, label, colour=TEAL):
    _rect(s, x, y, w, 1.02, CARD)
    _txt(s, x + 0.12, y + 0.13, w - 0.24, 0.42, value, 21, True, colour)
    _txt(s, x + 0.12, y + 0.60, w - 0.24, 0.34, label, 9.5, False, GREY)


def table(s, x, y, w, cols, rows, widths, head_fill=TEAL, rh=0.30):
    _rect(s, x, y, w, rh, head_fill, radius=False)
    cx = x
    for c, cw in zip(cols, widths):
        _txt(s, cx + 0.10, y + 0.06, cw - 0.16, 0.22, c, 9.5, True, WHITE)
        cx += cw
    for i, row in enumerate(rows):
        ry = y + rh + i * rh
        if i % 2 == 0:
            _rect(s, x, ry, w, rh, CARD, radius=False)
        cx = x
        for cell, cw in zip(row, widths):
            col, bold = BODY, False
            if isinstance(cell, tuple):
                cell, kind = cell
                col = {"good": GOOD, "bad": BAD, "teal": TEAL}.get(kind, BODY)
                bold = True
            _txt(s, cx + 0.10, ry + 0.06, cw - 0.16, 0.22, cell, 9.5, bold, col)
            cx += cw
    return y + rh + len(rows) * rh


# --------------------------------------------------------------- 01 title --
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = TEAL_DK
_rect(s, 0, 0, 0.09, 5.625, TEAL, radius=False)
_txt(s, 0.55, 0.50, 9.0, 0.30,
     "NTUT AIOT LAB  \u00b7  PROGRESS REPORT  \u00b7  AUGUST 2026", 10, True, LIGHT)
_txt(s, 0.55, 1.05, 8.9, 1.9,
     "Following a Named Object\nwith Vision and Language", 34, True, WHITE, spacing=0.95)
_txt(s, 0.55, 3.00, 8.6, 0.75,
     "You say what to follow. The drone finds it and chases it.\n"
     "The guardrail decides where it may not go.", 14, False, LIGHT)
_rect(s, 0.55, 4.10, 4.6, 0.02, TEAL, radius=False)
_txt(s, 0.55, 4.25, 8.0, 1.0,
     "Department of International Graduate Program, EECS\n"
     "Advisor: Prof. Kuan-Ting Lai   \u00b7   Presenter: Nathanael Cayadi",
     11, False, LIGHT)

# ------------------------------------------------------- 02 what changed --
s, y = slide("Headline", "Last report's pilot could not do this. This one can.",
             "Same guardrail, same simulator, same policy files \u2014 the pilot changed.")
table(s, 0.55, y, 8.9, ["", "Before  (AerialVLA 7B)", "Now  (this report)"], [
    ["Follow a named moving object", ("no \u2014 never moved", "bad"), ("yes", "good")],
    ["Language actually steers", ("no", "bad"), ("yes", "good")],
    ["Decision rate", ("0.11 Hz", "bad"), ("4\u20135 Hz", "good")],
    ["Altitude band held", ("escaped 31.6 s", "bad"), ("0.0 s", "good")],
    ["No-fly zone", "0.0 s", ("0.0 s", "good")],
], [3.9, 2.5, 2.5])
_txt(s, 0.55, 4.95, 8.9, 0.4,
     "Source: docs/FINDING-what-drives-aerialvla.md, docs/RESULT-vlm-follow.md",
     9, False, GREY)

# ------------------------------------------------------ 03 what we found --
s, y = slide("The finding", "AerialVLA's language input never reached its output.",
             "108 controlled forward passes per adapter, on real captured frames.")
kpi(s, 0.55, y, 2.85, "12 / 12", "flights: identical 'stop and land' output")
kpi(s, 3.58, y, 2.85, "0.05 vs 0.08", "yaw for correct vs WRONG colour word")
kpi(s, 6.61, y, 2.84, "11 / 18", "frames answered LAND with no compass hint")
_rect(s, 0.55, y + 1.25, 8.9, 1.55, CARD)
_txt(s, 0.80, y + 1.42, 8.4, 0.32,
     "The prompt has two slots. Only one of them steers.", 13, True, INK)
_txt(s, 0.80, y + 1.82, 8.4, 1.1,
     "\"Fly {direction} and find the target. {object}\"\n"
     "{direction} is a compass phrase computed from ground-truth coordinates \u2014 it orders the commanded\n"
     "yaw monotonically from \u22120.47 rad/s (\"to your left\") to +0.16 (\"to your right\").\n"
     "{object} is the language. Correct and wrong colour words gave indistinguishable actions.",
     10.5, False, GREY)

# ------------------------------------------------------------- 04 the fix --
s, y = slide("The fix", "We changed the model, not the goal.",
             "Still vision \u2192 language \u2192 action. Now with a model whose words reach the output.")
for i, (head, body) in enumerate([
    ("OWL-ViT  153 M", "Open-vocabulary detector.\nText in, boxes out.\n34\u201356 ms per frame."),
    ("Servo controller", "Box offset \u2192 yaw.\nBox width \u2192 speed.\nAltitude error \u2192 climb."),
    ("Guardrail \u2014 unchanged", "Same Shield, same policies.\nNo modification needed to\naccept a new kind of pilot."),
]):
    x = 0.55 + i * 3.06
    _rect(s, x, y, 2.78, 1.75, CARD)
    _rect(s, x, y, 2.78, 0.06, TEAL, radius=False)
    _txt(s, x + 0.18, y + 0.24, 2.42, 0.3, head, 12.5, True, INK)
    _txt(s, x + 0.18, y + 0.62, 2.42, 1.0, body, 10.5, False, GREY)
_txt(s, 0.55, y + 2.00, 8.9, 0.7,
     "Nothing in the steering path knows where the car is. Its position is used only to spawn it,\n"
     "and offline, to score the flight.", 11, False, BODY)

# ----------------------------------------------------------- 05 the demo --
s, y = slide("The demo", "Two flights, one command.", ".\\run_follow_vlm.ps1")
_rect(s, 0.55, y, 4.35, 1.9, CARD)
_txt(s, 0.75, y + 0.20, 3.95, 0.3, "1 \u00b7 Tracking, with stops", 13, True, INK)
_txt(s, 0.75, y + 0.60, 3.95, 1.2,
     "The car drives, stops for 6 s, drives,\nstops, drives again. The drone has to\n"
     "stop too \u2014 that is the proof.", 10.5, False, GREY)
_rect(s, 5.10, y, 4.35, 1.9, CARD)
_txt(s, 5.30, y + 0.20, 3.95, 0.3, "2 \u00b7 The same, with a no-fly zone", 13, True, INK)
_txt(s, 5.30, y + 0.60, 3.95, 1.2,
     "A fence spans the corridor. The car\ndrives through it. The drone brakes,\n"
     "holds at 3 m, and does not.", 10.5, False, GREY)
_txt(s, 0.55, y + 2.10, 8.9, 0.7,
     "Recorded from two cameras: the drone's own view with the detection box and live telemetry,\n"
     "and a third-person chase view, stitched side by side.", 11, False, BODY)

# ------------------------------------------------------------ 06 results --
s, y = slide("Results", "It follows, and the rules hold.",
             "Car at 2.0 m/s. Separation measured against its true position, offline.")
table(s, 0.55, y, 8.9,
      ["Condition", "Detector hit", "Mean sep.", "Within 30 m", "Shield overrides", "NFZ time"], [
    ["Follow: car stops twice", "100%", "15.6 m", ("100%", "good"), "0", ("0.0 s", "good")],
    ["Follow + no-fly zone", "67.8%", "45.2 m", ("28.6%", "bad"), ("0", "teal"), ("0.0 s", "good")],
    ["Control: wrong colour", "99.0%", ("68.5 m", "bad"), ("20.7%", "bad"), "27", "0.0 s"],
], [2.6, 1.30, 1.20, 1.25, 1.35, 1.20])
_txt(s, 0.55, y + 1.35, 8.9, 1.25,
     "Row 1: the drone stops when the car stops \u2014 1.17 m/s while the car moves against 0.36 m/s while it is\n"
     "parked, correlation +0.55. A follower has to stop, so this is the evidence that it tracks the car\n"
     "rather than just flying down the same street.\n"
     "Row 2 is a deliberate failure: the fence has no way round, so the car escapes and tracking collapses\n"
     "to 28.6% \u2014 with zero Shield overrides, because the controller now stops itself, and the zone still\n"
     "never entered.", 10, False, GREY)

# ------------------------------------------------------ 07 it's the words --
s, y = slide("Proof", "One word decides whether it follows.",
             "Identical scene, identical orange car, identical settings.")
_rect(s, 0.55, y, 4.35, 1.85, CARD)
_txt(s, 0.75, y + 0.22, 3.95, 0.34, "\"an orange car\"", 15, True, GOOD)
_txt(s, 0.75, y + 0.70, 3.95, 0.9,
     "100% of the flight within 30 m\nmean separation 15.6 m", 12, False, BODY)
_rect(s, 5.10, y, 4.35, 1.85, CARD)
_txt(s, 5.30, y + 0.22, 3.95, 0.34, "\"a blue car\"", 15, True, BAD)
_txt(s, 5.30, y + 0.70, 3.95, 0.9,
     "20.7% within 30 m\nmean separation 68.5 m", 12, False, BODY)
_txt(s, 0.55, y + 2.05, 8.9, 0.8,
     "This is the control that matters. Without it, \"the drone ended up near the car\" could just be drift.\n"
     "The colour word is verified against the pixels in the box \u2014 the detector grounds the noun,\n"
     "an HSV check grounds the adjective.", 11, False, BODY)

# ----------------------------------------------------- 08 guardrail works --
s, y = slide("Safety", "Safe is not enough \u2014 it also has to look safe.",
             "The no-fly-zone flight, before and after the controller learned about fences.")
kpi(s, 0.55, y, 2.85, "659 \u2192 0", "Shield overrides, after the fix", GOOD)
kpi(s, 3.58, y, 2.85, "3.00 m", "closest approach; stand-off is 3 m", TEAL)
kpi(s, 6.61, y, 2.84, "0.0 s", "inside the zone, every flight", GOOD)
_rect(s, 0.55, y + 1.25, 8.9, 1.5, CARD)
_txt(s, 0.80, y + 1.42, 8.4, 0.3,
     "Two layers arguing ten times a second looks exactly as unsafe as it is.", 12.5, True, INK)
_txt(s, 0.80, y + 1.80, 8.4, 0.95,
     "The servo asked to close on the car every tick and the Shield refused every tick. Neither changed its\n"
     "mind, so the aircraft chattered against the boundary \u2014 659 corrections in 795 ticks. The controller\n"
     "now knows the geofence itself (mission data, not target data), brakes from 12 m and holds at 3 m,\n"
     "while yaw keeps tracking so the target stays in view. HUD reads NFZ AHEAD \u2014 HOLDING.",
     10.5, False, GREY)

# --------------------------------------------- 09 how the layers connect --
s, y = slide("Q: how do they connect?",
             "The pilot and the guardrail share one struct.",
             "That contract is why the pilot could be replaced without touching the safety layer.")
_rect(s, 0.55, y, 4.30, 1.30, CARD)
_txt(s, 0.75, y + 0.14, 3.9, 0.24, "guardrail/models.py", 10, True, TEAL)
_txt(s, 0.75, y + 0.44, 3.95, 0.8,
     "class Action4D:\n"
     "    vx, vy      # m/s, North / East\n"
     "    vz_up       # m/s, up\n"
     "    yaw_rate    # +clockwise", 10, False, BODY, spacing=1.05)
_rect(s, 5.15, y, 4.30, 1.30, CARD)
_txt(s, 5.35, y + 0.14, 3.9, 0.24, "every tick, in the flight loop", 10, True, TEAL)
_txt(s, 5.35, y + 0.44, 3.95, 0.8,
     "raw    = Action4D(...)     # pilot asks\n"
     "smooth = limiter(raw)\n"
     "d      = shield.filter(state, smooth)\n"
     "drone.move_by_velocity(d.emitted)", 10, False, BODY, spacing=1.05)
_txt(s, 0.55, y + 1.44, 8.9, 0.3,
     "`raw` never reaches the simulator. Only `d.emitted` flies \u2014 the guardrail is the last authority.",
     11, True, INK)
table(s, 0.55, y + 1.80, 8.9,
      ["Case", "Pilot asked for", "What actually flew", "Rule at risk"], [
    ["Guardrail agrees", "vx +0.41  vy \u22120.62", "vx +0.40  vy \u22120.40", ("none", "good")],
    ["Controller brakes itself", "vy +0.74", "vy +0.74 (unchanged)", ("none \u2014 fence 6.0 m", "good")],
    ["Guardrail overrides", ("vx \u22120.47 toward a wall", "bad"), ("vx +0.50", "teal"),
     ("clearance 4.98 < 5.0 m", "bad")],
], [2.2, 2.30, 2.25, 2.15], rh=0.32)
_txt(s, 0.55, y + 2.90, 8.9, 0.3,
     "Real ticks from demo/out/*/flight_log.jsonl. yaw_rate is identical in all three \u2014 the Shield never "
     "touches heading.", 9.5, False, GREY)

# -------------------------------------------------- 10 detect and turn --
s, y = slide("Q: how does it detect and turn?",
             "From a word to a yaw command, in four steps.")
for head, body in [
    ("1 \u00b7 Detect",
     "OWL-ViT takes the phrase and returns scored boxes. 34\u201356 ms, in its own thread."),
    ("2 \u00b7 Verify the colour",
     "HSV check inside each box. The detector grounds the noun, this grounds the adjective.\n"
     "Winner is score \u00d7 (0.25 + 0.75 \u00d7 colour) \u2014 not score alone, which once locked onto a building."),
    ("3 \u00b7 Turn",
     "off = (cx \u2212 W/2) / (W/2)    bearing = 45\u00b0 \u00d7 off    yaw_rate = gain \u00d7 bearing\n"
     "The pixel offset becomes a REAL angle, because it is scaled by the camera's half-FOV."),
    ("4 \u00b7 Close or back off",
     "Box width stands in for distance: too small \u2192 forward, too large \u2192 reverse.\n"
     "Scaled by cos(bearing), so it never charges while the target is off to one side."),
]:
    _rect(s, 0.55, y, 0.06, 0.80, TEAL, radius=False)
    _txt(s, 0.80, y + 0.01, 8.6, 0.28, head, 12, True, INK)
    _txt(s, 0.80, y + 0.32, 8.6, 0.5, body, 10, False, GREY)
    y += 0.88
_rect(s, 0.55, y + 0.04, 8.9, 0.60, CARD)
_txt(s, 0.78, y + 0.15, 8.4, 0.42,
     "When detection drops: COAST on the last motion (< 2 s), then SEARCH toward the side it was last seen,\n"
     "then SCAN slowly. It never freezes, and never-yet-acquired counts as SEARCH rather than give-up.",
     10, False, BODY)

# ------------------------------------------------------------- 11 limits --
s, y = slide("Limits", "What this does not do.",
             "Stated here so nobody discovers it during questions.")
for head, body in [
    ("The noun does most of the work",
     "\"a car\" alone chases city clutter and ends 51 m away. The colour check helps, but this grounds\n"
     "\"car\", not \"that particular one\"."),
    ("It cannot say the object is absent",
     "With no car in the scene the raw detector still fires on ~75% of frames. Only the colour check\n"
     "suppresses it."),
    ("The target is made easy, and it is a simulator",
     "The car drives straight and parks, because turns swing it through the blind spot. The obstacle map\n"
     "holds buildings only \u2014 not trees or lamp posts. One flight per condition, and no hardware yet."),
]:
    _rect(s, 0.55, y, 8.9, 0.95, CARD)
    _txt(s, 0.78, y + 0.13, 8.4, 0.28, head, 12, True, INK)
    _txt(s, 0.78, y + 0.45, 8.4, 0.45, body, 10, False, GREY)
    y += 1.08

# --------------------------------------------------------------- 12 next --
s, y = slide("Next", "Where this goes.")
for head, body in [
    ("Confidence, not just position",
     "A track-confirmation stage, so the system can report that the target is gone instead of\n"
     "locking onto whatever is nearest."),
    ("A fence with a gap",
     "Built and unit-tested \u2014 it picks the open side and refuses when there is none \u2014 but not yet\n"
     "demonstrated in flight: the aircraft met street furniture the obstacle map does not contain."),
    ("Toward hardware",
     "The guardrail already runs on any action source. The open questions are latency and camera\n"
     "geometry on a real airframe."),
]:
    _rect(s, 0.55, y, 0.06, 0.85, TEAL, radius=False)
    _txt(s, 0.80, y + 0.02, 8.6, 0.3, head, 12.5, True, INK)
    _txt(s, 0.80, y + 0.36, 8.6, 0.5, body, 10.5, False, GREY)
    y += 1.05

# ----------------------------------------------------------- 13 takeaway --
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = TEAL_DK
_rect(s, 0, 0, 0.09, 5.625, TEAL, radius=False)
_txt(s, 0.55, 0.60, 9.0, 0.30, "TAKEAWAY", 10, True, LIGHT)
_txt(s, 0.55, 1.15, 8.9, 1.7,
     "The pilot changed completely.\nThe safety layer did not have to.", 30, True, WHITE, spacing=0.95)
_txt(s, 0.55, 3.05, 8.7, 1.3,
     "A 7B action model was swapped for a 153 M detector and a servo loop \u2014 a different kind of\n"
     "controller entirely \u2014 and the Shield accepted it without a single change to its code or its\n"
     "policy files. Across every flight in this report: no-fly zone 0.0 s, altitude escape 0.0 s.\n"
     "That portability is the architectural claim, and it survived replacing the brain.",
     12.5, False, LIGHT)
_rect(s, 0.55, 4.60, 4.6, 0.02, TEAL, radius=False)
_txt(s, 0.55, 4.75, 8.9, 0.4,
     "Reproduce: .\\run_follow_vlm.ps1   \u00b7   TUTORIAL-DEMO-FOLLOW.md", 11, False, LIGHT)

OUT.parent.mkdir(parents=True, exist_ok=True)
try:
    prs.save(OUT)
    dest = OUT
except PermissionError:
    # PowerPoint holds an exclusive lock on an open file. Write beside it rather
    # than failing, so a rebuild never loses work because a window was left open.
    dest = OUT.with_name(OUT.stem + "-new" + OUT.suffix)
    prs.save(dest)
    print(f"[deck] {OUT.name} is open in PowerPoint - wrote {dest.name} instead")
print(f"[deck] {len(prs.slides._sldIdLst)} slides -> {dest}")
